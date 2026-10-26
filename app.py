import hashlib
import json
import os
import re
import sqlite3
import uuid
import io
from datetime import datetime, date
from pathlib import Path

import streamlit as st

# Optional local OCR support. Tesseract must be installed on the computer.
TESSERACT = Path(r"C:\Program Files\Tesseract-OCR")
TESSDATA = TESSERACT / "tessdata"
if TESSERACT.is_dir():
    os.environ["PATH"] = str(TESSERACT) + os.pathsep + os.environ.get("PATH", "")

import pymupdf

ROOT = Path(__file__).parent
RULES_FILE = ROOT / "compliance_rules.json"
RECORDS_FILE = ROOT / "demo_institution_records.json"
DB_FILE = ROOT / "ai_verify.db"

FIELDS = {
    "Institution name": r"(?im)^\s*(?:institution\s+name|name\s+of\s+institution)\s*[:\-]\s*(.+?)\s*$",
    "Institute ID": r"(?im)^\s*(?:institute\s+id|institute\s+pid|pid|aishe\s+code)\s*[:\-]\s*(.+?)\s*$",
    "Registration number": r"(?im)^\s*(?:registration\s+number|registration\s+no\.?)\s*[:\-]\s*(.+?)\s*$",
    "Course": r"(?im)^\s*(?:course|program|programme)\s*[:\-]\s*(.+?)\s*$",
    "Approved intake": r"(?im)^\s*(?:approved\s+intake|intake)\s*[:\-]\s*(.+?)\s*$",
    "Faculty count": r"(?im)^\s*(?:faculty\s+count|number\s+of\s+faculty|faculty)\s*[:\-]\s*(.+?)\s*$",
    "Issue date": r"(?im)^\s*(?:issue\s+date|date\s+of\s+issue)\s*[:\-]\s*(.+?)\s*$",
    "Expiry date": r"(?im)^\s*(?:expiry\s+date|expiration\s+date|valid\s+until|validity\s+date)\s*[:\-]\s*(.+?)\s*$",
    "Academic year": r"(?im)^\s*(?:academic\s+year|session)\s*[:\-]\s*(.+?)\s*$",
    "Accreditation": r"(?im)^\s*(?:accreditation|accreditation\s+status)\s*[:\-]\s*(.+?)\s*$",
    "NOC number": r"(?im)^\s*(?:noc\s+number|noc\s+no\.?)\s*[:\-]\s*(.+?)\s*$",
    "TER payment status": r"(?im)^\s*(?:ter\s+payment\s+status|payment\s+status)\s*[:\-]\s*(.+?)\s*$",
    "Application reference": r"(?im)^\s*(?:application\s+reference|application\s+id|application\s+number)\s*[:\-]\s*(.+?)\s*$",
    "Address": r"(?im)^\s*(?:institution\s+address|address)\s*[:\-]\s*(.+?)\s*$",
    "Built-up area": r"(?im)^\s*(?:built[ -]?up\s+area|builtup\s+area)\s*[:\-]\s*(.+?)\s*$",
    "Laboratory count": r"(?im)^\s*(?:laboratory\s+count|number\s+of\s+laboratories|labs)\s*[:\-]\s*(.+?)\s*$",
    "Library books": r"(?im)^\s*(?:library\s+books|number\s+of\s+books)\s*[:\-]\s*(.+?)\s*$",
    "Hostel capacity": r"(?im)^\s*(?:hostel\s+capacity|hostel\s+seats)\s*[:\-]\s*(.+?)\s*$",
    "Document reference number": r"(?im)^\s*(?:document\s+reference|certificate\s+number|reference\s+number)\s*[:\-]\s*(.+?)\s*$",
}

st.set_page_config(
    page_title="AI VERIFY | Review Workspace",
    page_icon="📄",
    layout="wide",
)

st.markdown("""
<style>
.stApp,
[data-testid="stAppViewContainer"] {
    background: #f3f7fc !important;
    color: #17243b !important;
}

.block-container {
    max-width: 1450px;
    padding: 2rem 2.4rem 4rem;
}

[data-testid="stSidebar"] {
    background: #eaf1fa !important;
    border-right: 1px solid #d4deeb;
}

[data-testid="stSidebar"] * {
    color: #17243b !important;
}

h1, h2, h3, p, label,
[data-testid="stWidgetLabel"] p {
    color: #17243b !important;
}

/* Keep the existing colorful header */
.hero {
    padding: 30px 34px;
    border: 1px solid #ffffff24;
    border-radius: 24px;
    margin-bottom: 18px;
    background: linear-gradient(115deg, #1e3a68, #2863a8 58%, #0e8c91);
    box-shadow: 0 20px 45px #1e3a6822;
}

.hero h1,
.hero p {
    color: #ffffff !important;
}

.notice {
    padding: 14px 18px;
    border: 1px solid #c8e2ef;
    border-left: 4px solid #08a6a0;
    border-radius: 13px;
    background: #eef8fc;
    color: #173657;
    margin: 12px 0 22px;
}

/* Make every text box visibly outlined before it is clicked */
[data-testid="stTextInput"] [data-baseweb="input"] {
    background: #ffffff !important;
    border: 1px solid #9aacc2 !important;
    border-radius: 9px !important;
    box-shadow: inset 0 1px 2px #17243b0a !important;
}

[data-testid="stTextInput"] input {
    background: transparent !important;
    color: #17243b !important;
    border: 0 !important;
    caret-color: #17243b !important;
}

[data-testid="stTextInput"] [data-baseweb="input"]:focus-within {
    border: 1px solid #315ee7 !important;
    box-shadow: 0 0 0 3px #315ee722 !important;
}

/* Make text areas visibly outlined too */
[data-testid="stTextArea"] [data-baseweb="textarea"] {
    background: #ffffff !important;
    border: 1px solid #9aacc2 !important;
    border-radius: 9px !important;
}

[data-testid="stTextArea"] textarea {
    background: transparent !important;
    color: #17243b !important;
    border: 0 !important;
}

[data-testid="stTextArea"] [data-baseweb="textarea"]:focus-within {
    border: 1px solid #315ee7 !important;
    box-shadow: 0 0 0 3px #315ee722 !important;
}

/* Make dropdowns distinct from the page */
[data-testid="stSelectbox"] [data-baseweb="select"] > div {
    background: #ffffff !important;
    color: #17243b !important;
    border: 1px solid #9aacc2 !important;
    border-radius: 9px !important;
}

/* White cards and upload area */
[data-testid="stMetric"] {
    padding: 18px;
    border: 1px solid #dce5ef;
    border-radius: 17px;
    background: #ffffff;
    box-shadow: 0 8px 22px #243b5310;
}

[data-testid="stMetricValue"] {
    color: #315ee7 !important;
}

[data-testid="stMetricLabel"] {
    color: #52657d !important;
}

[data-testid="stFileUploader"] {
    padding: 10px;
    border: 1px dashed #9aacc2;
    border-radius: 16px;
    background: #ffffff;
}

/* Buttons */
.stButton > button,
.stDownloadButton > button,
[data-testid="stFormSubmitButton"] > button {
    border: 0;
    border-radius: 11px;
    color: #ffffff !important;
    font-weight: 700;
    background: linear-gradient(100deg, #315ee7, #2386b8);
}

.stButton > button:hover,
.stDownloadButton > button:hover {
    color: #ffffff !important;
    border: 0;
}

/* Expanders and tabs */
[data-testid="stExpander"] {
    border: 1px solid #c7d2e0;
    border-radius: 13px;
    background: #ffffff;
}

[data-testid="stTabs"] button[role="tab"] {
    color: #263b54 !important;
    font-weight: 700;
}
/* Strong outline for empty and filled text boxes */
[data-testid="stTextInput"] input,
[data-testid="stTextArea"] textarea {
    background: #ffffff !important;
    color: #17243b !important;
    border: 2px solid #64748b !important;
    border-radius: 8px !important;
    box-shadow: none !important;
    opacity: 1 !important;
}

[data-testid="stTextInput"] input:focus,
[data-testid="stTextArea"] textarea:focus {
    border: 2px solid #315ee7 !important;
    box-shadow: 0 0 0 3px #315ee733 !important;
}
</style>
""", unsafe_allow_html=True)
CSS_FILE = Path(__file__).with_name("style.css")
if CSS_FILE.is_file():
    st.markdown(
        f"<style>{CSS_FILE.read_text(encoding='utf-8')}</style>",
        unsafe_allow_html=True,
    )

def load_json(path, fallback):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return fallback


RULES = load_json(
    RULES_FILE,
    {"categories": {}, "notice": "Demo rules only."},
)
DEMO_RECORDS = load_json(
    RECORDS_FILE,
    {"records": []},
).get("records", [])


def setup_db():
    with sqlite3.connect(DB_FILE) as db:
        db.execute(
            "CREATE TABLE IF NOT EXISTS audit("
            "event_id TEXT PRIMARY KEY, event_time TEXT, user_role TEXT, "
            "event_type TEXT, details TEXT)"
        )
        db.execute(
            "CREATE TABLE IF NOT EXISTS reviews("
            "review_id INTEGER PRIMARY KEY AUTOINCREMENT, finding_id TEXT, "
            "action TEXT, comment TEXT, event_time TEXT)"
        )


def log_event(role, kind, details):
    with sqlite3.connect(DB_FILE) as db:
        db.execute(
            "INSERT INTO audit VALUES(?,?,?,?,?)",
            (
                str(uuid.uuid4()),
                datetime.now().astimezone().isoformat(timespec="seconds"),
                role,
                kind,
                details,
            ),
        )


def pdf_text(pdf_bytes):
    pages, metadata, problems = [], {}, []
    try:
        with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
            metadata = pdf.metadata or {}
            for number, page in enumerate(pdf, 1):
                text = page.get_text("text").strip()
                if not text:
                    try:
                        text_page = page.get_textpage_ocr(
                            language="eng",
                            dpi=300,
                            full=True,
                            tessdata=str(TESSDATA),
                        )
                        text = page.get_text(
                            "text",
                            textpage=text_page,
                        ).strip()
                        if not text:
                            problems.append(
                                f"OCR found no readable text on page {number}"
                            )
                    except Exception as exc:
                        problems.append(
                            f"OCR failed on page {number}: {exc}"
                        )
                pages.append(text)
        return pages, metadata, problems
    except Exception as exc:
        return [], {}, [f"PDF could not be opened: {exc}"]


def extract_file(file_name, data):
    """Extract text where practical. Video/photo evidence remains human-reviewed."""
    suffix = Path(file_name).suffix.casefold()
    metadata, problems, pages = {}, [], []

    try:
        if suffix == ".pdf":
            return pdf_text(data)

        if suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            chunks = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for table in document.tables:
                chunks.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                )
            pages = ["\n".join(chunks)]

        elif suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=True,
            )
            for sheet in workbook.worksheets:
                rows = [
                    " | ".join(str(value) for value in row if value is not None)
                    for row in sheet.iter_rows(values_only=True)
                ]
                pages.append(
                    f"Sheet: {sheet.title}\n" +
                    "\n".join(row for row in rows if row)
                )

        elif suffix == ".csv":
            pages = [data.decode("utf-8-sig", errors="replace")]

        elif suffix in (".txt", ".md"):
            pages = [data.decode("utf-8-sig", errors="replace")]

        elif suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            try:
                from PIL import Image
                import pytesseract

                image = Image.open(io.BytesIO(data))
                pages = [pytesseract.image_to_string(image)]
                if not pages[0].strip():
                    problems.append(
                        "No readable text detected in image; inspect it manually."
                    )
            except Exception as exc:
                problems.append(
                    f"Image retained for human review; OCR unavailable: {exc}"
                )

        elif suffix in (".mp4", ".mov", ".avi", ".webm"):
            problems.append(
                "Video evidence is available for this session only. "
                "No video analysis is configured; review it manually."
            )
            metadata["evidence_type"] = "video"

        else:
            problems.append(
                "Unsupported format for text extraction; review the evidence manually."
            )

    except Exception as exc:
        problems.append(f"Could not read this file: {exc}")

    if pages and not any(page.strip() for page in pages):
        problems.append(
            "No selectable text found. The file may be scanned or image-only; "
            "OCR or manual review is needed."
        )

    return pages, metadata, problems


def parse_fields(pages):
    text = "\n".join(pages)
    values, evidence = {}, {}

    for label, pattern in FIELDS.items():
        match = re.search(pattern, text)
        values[label] = match.group(1).strip() if match else ""

        if match:
            offset = match.start()
            line_number = text[:offset].count("\n") + 1
            evidence[label] = (
                f"approx. extracted-text line {line_number} "
                "(not a PDF page number)"
            )
        else:
            evidence[label] = "not found"

    return values, evidence


def norm(value):
    return re.sub(r"[^a-z0-9]", "", str(value).casefold())


def parse_date(value):
    value = str(value).strip()
    formats = (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%m/%d/%Y",
    )
    for date_format in formats:
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            pass
    return None


def add(findings, status, source, reason, evidence=""):
    findings.append({
        "id": str(uuid.uuid4()),
        "status": status,
        "source": source,
        "reason": reason,
        "evidence": evidence,
    })


def run_checks(docs, category, references, role):
    findings = []
    category_rules = RULES.get("categories", {}).get(category, {})
    required_fields = category_rules.get("required_fields", [])
    required_docs = category_rules.get("required_documents", [])
    supplied_doc_types = {doc["document_type"] for doc in docs}

    for required_doc in required_docs:
        if required_doc not in supplied_doc_types:
            add(
                findings,
                "Missing document",
                "Checklist",
                f"Required document type not uploaded: {required_doc}.",
            )

    all_fields = {}
    hashes = {}

    for doc in docs:
        for label in FIELDS:
            value = doc["fields"].get(label, "").strip()
            if value:
                all_fields.setdefault(label, []).append(
                    (
                        doc["name"],
                        value,
                        doc["evidence"].get(label, ""),
                    )
                )

        digest = doc["hash"]
        if digest in hashes:
            add(
                findings,
                "Duplicate indicator",
                doc["name"],
                f"Exact file contents match {hashes[digest]}.",
            )
        else:
            hashes[digest] = doc["name"]

        expiry = doc["fields"].get("Expiry date", "")
        if expiry:
            parsed_expiry = parse_date(expiry)
            if parsed_expiry and parsed_expiry < date.today():
                add(
                    findings,
                    "Expired — review",
                    doc["name"],
                    f"Expiry date {expiry} is before today.",
                    doc["evidence"].get("Expiry date", ""),
                )
            elif not parsed_expiry:
                add(
                    findings,
                    "Date needs review",
                    doc["name"],
                    f"Could not read expiry date format: {expiry}.",
                )

        issue = doc["fields"].get("Issue date", "")
        parsed_issue = parse_date(issue) if issue else None
        if parsed_issue and parsed_issue > date.today():
            add(
                findings,
                "Date anomaly — review",
                doc["name"],
                f"Issue date {issue} is in the future.",
                doc["evidence"].get("Issue date", ""),
            )

        for problem in doc["problems"]:
            add(findings, "OCR / file review", doc["name"], problem)

        created = doc["metadata"].get("creationDate", "")
        modified = doc["metadata"].get("modDate", "")
        if created and modified and created != modified:
            add(
                findings,
                "Metadata indicator — review",
                doc["name"],
                "PDF creation and modification timestamps differ; "
                "ordinary editing can cause this.",
            )

    for label in required_fields:
        if not all_fields.get(label):
            add(
                findings,
                "Missing information",
                "Application package",
                f"Required field not found in the uploaded package: {label}.",
            )

    for label, entries in all_fields.items():
        if len({norm(value) for _, value, _ in entries}) > 1:
            detail = "; ".join(
                f"{name} ({evidence}): {value}"
                for name, value, evidence in entries
            )
            add(
                findings,
                "Cross-document mismatch — review",
                "Across uploaded documents",
                f"{label} differs: {detail}",
                detail,
            )

    # This is a fictional local registry, not an AICTE connection.
    institute_id = next(
        (value for _, value, _ in all_fields.get("Institute ID", [])),
        "",
    )
    record = next(
        (
            row for row in DEMO_RECORDS
            if norm(row.get("institute_id")) == norm(institute_id)
        ),
        None,
    )

    if institute_id and record:
        comparisons = (
            ("Institution name", "institution_name"),
            ("Registration number", "registration_number"),
            ("Academic year", "academic_year"),
        )
        for field, record_key in comparisons:
            submitted = next(
                (value for _, value, _ in all_fields.get(field, [])),
                "",
            )
            demo_value = record.get(record_key, "")
            if submitted and demo_value:
                status = (
                    "Synthetic registry match"
                    if norm(submitted) == norm(demo_value)
                    else "Synthetic registry mismatch — review"
                )
                add(
                    findings,
                    status,
                    "Bundled synthetic demo registry",
                    f"{field}: submitted “{submitted}”; "
                    f"demo record “{demo_value}”.",
                )

        submitted_course = next(
            (value for _, value, _ in all_fields.get("Course", [])),
            "",
        )
        approved_courses = record.get("approved_courses", [])
        if submitted_course and approved_courses:
            matched = any(
                norm(submitted_course) == norm(course)
                for course in approved_courses
            )
            add(
                findings,
                "Synthetic registry match"
                if matched
                else "Synthetic registry mismatch — review",
                "Bundled synthetic demo registry",
                f"Course: submitted “{submitted_course}”; "
                f"demo record lists {', '.join(approved_courses)}.",
            )

        submitted_intake = next(
            (value for _, value, _ in all_fields.get("Approved intake", [])),
            "",
        )
        trusted_intake = record.get("approved_intake")
        if submitted_intake and trusted_intake is not None:
            matched = norm(submitted_intake) == norm(trusted_intake)
            add(
                findings,
                "Synthetic registry match"
                if matched
                else "Synthetic registry mismatch — review",
                "Bundled synthetic demo registry",
                f"Approved intake: submitted “{submitted_intake}”; "
                f"demo record “{trusted_intake}”.",
            )

        add(
            findings,
            "Synthetic record found",
            "Bundled synthetic demo registry",
            f"Demo-only record: {record.get('approval_status', 'no status')} "
            f"— {record.get('academic_year', '')}. Not an AICTE response.",
        )

    elif institute_id:
        add(
            findings,
            "Live source unavailable — manual review",
            "AICTE source",
            "No authorized AICTE API is configured. The app did not query AICTE.",
        )

    for field, reference in references.items():
        if not reference.strip():
            continue

        submitted_values = all_fields.get(field, [])
        if not submitted_values:
            add(
                findings,
                "Reference check — value missing",
                field,
                f"Entered reference exists, but no {field} was extracted.",
            )

        for filename, value, evidence in submitted_values:
            status = (
                "Entered reference matches"
                if norm(value) == norm(reference)
                else "Entered reference mismatch — review"
            )
            add(
                findings,
                status,
                filename,
                f"{field}: file text “{value}”; "
                f"reviewer-entered reference “{reference}”. "
                "This is not an external verification.",
                evidence,
            )

    if not findings:
        add(
            findings,
            "No demo issues found",
            "Uploaded documents",
            "No configured demo checks raised an issue. "
            "This does not establish authenticity or regulatory compliance.",
        )

    log_event(
        role,
        "analysis",
        f"category={category}; files={len(docs)}; findings={len(findings)}",
    )
    log_event(
        role,
        "rules_applied",
        json.dumps(category_rules, ensure_ascii=True),
    )
    for item in findings:
        log_event(
            role,
            "finding",
            f"{item['status']} | {item['source']} | {item['reason']} "
            f"| evidence={item.get('evidence', '')}",
        )

    return findings


def make_report(findings, category, app_label):
    pdf = pymupdf.open()
    page = pdf.new_page()
    y = 55

    def line(text, size=10, bold=False):
        nonlocal page, y
        text = str(text).encode("ascii", "replace").decode("ascii")
        chunks = [
            text[index:index + 110]
            for index in range(0, max(len(text), 1), 110)
        ]
        for chunk in chunks:
            if y > 775:
                page = pdf.new_page()
                y = 55
            page.insert_text(
                (48, y),
                chunk,
                fontsize=size,
                fontname="hebo" if bold else "helv",
            )
            y += 16

    line("AI VERIFY - APPLICATION REVIEW REPORT", 17, True)
    line(f"Application: {app_label}")
    line(f"Category: {category}")
    line(
        "Generated: " +
        datetime.now().astimezone().isoformat(timespec="seconds")
    )
    line("Prototype only. No live AICTE system was queried.", 11, True)
    line("")

    for index, item in enumerate(findings, 1):
        line(f"{index}. {item['status']} | {item['source']}", 11, True)
        line(item["reason"])
        if item.get("evidence"):
            line(f"Evidence: {item['evidence']}")
        line("")

    line("All flags are advisory. Final decisions belong to authorized officials.")
    return pdf.tobytes()


def status_counts(findings):
    flagged = sum(
        any(
            word in finding["status"].casefold()
            for word in (
                "mismatch",
                "missing",
                "expired",
                "anomaly",
                "review",
                "unavailable",
                "indicator",
            )
        )
        for finding in findings
    )
    passed = sum(
        "match" in finding["status"].casefold()
        or "no demo issues" in finding["status"].casefold()
        for finding in findings
    )
    pending = sum(
        "pending" in finding["status"].casefold()
        or "unavailable" in finding["status"].casefold()
        for finding in findings
    )
    return passed, flagged, pending


def show_audit():
    with sqlite3.connect(DB_FILE) as db:
        return db.execute(
            "SELECT event_time,user_role,event_type,details "
            "FROM audit ORDER BY rowid DESC LIMIT 100"
        ).fetchall()


setup_db()

with st.sidebar:
    st.markdown("### Workspace")
    role = st.selectbox(
        "Demo role",
        ["Applicant", "Reviewer", "Administrator"],
    )
    st.caption("Role selector is for demonstration; it is not real access control.")

    categories = list(RULES.get("categories", {}).keys()) or [
        "Extension of Approval (EoA)"
    ]
    category = st.selectbox("Application category", categories)

    st.markdown("### Reference values (optional)")
    reference_fields = [
        "Institution name",
        "Institute ID",
        "Registration number",
        "Course",
        "Approved intake",
        "Academic year",
        "Faculty count",
        "Accreditation",
        "NOC number",
        "TER payment status",
        "Application reference",
        "Address",
        "Built-up area",
        "Laboratory count",
        "Library books",
        "Hostel capacity",
        "Document reference number",
    ]
    references = {
        field: st.text_input(field, key="ref_" + field)
        for field in reference_fields
    }

    app_label = st.text_input("Application label", "Demo EoA application")
    st.markdown("---")
    st.caption("Synthetic records are fictional. No live AICTE connection is configured.")

st.markdown(
    '<div class="hero"><h1>📄 AI VERIFY</h1>'
    '<p>Application scrutiny workspace · Evidence-led · Human-reviewed</p></div>',
    unsafe_allow_html=True,
)
st.markdown(
    '<div class="notice">AICTE LIVE STATUS: Not connected. This prototype '
    'uses configured demo rules, reviewer-entered values, and a clearly '
    'labeled synthetic registry. It never treats an alert as proof of fraud.</div>',
    unsafe_allow_html=True,
)

tab_upload, tab_dashboard, tab_audit = st.tabs([
    "📤 Application review",
    "📊 Dashboard & report",
    "🧾 Audit trail",
])

with tab_upload:
    st.subheader("1 · Upload application documents")

    accepted_types = [
        "pdf", "docx", "xlsx", "xlsm", "csv", "txt", "md",
        "png", "jpg", "jpeg", "tif", "tiff",
        "mp4", "mov", "avi", "webm",
    ]
    uploaded = st.file_uploader(
        "Add application documents and evidence",
        type=accepted_types,
        accept_multiple_files=True,
        help=(
            "Text extraction is supported for PDF, DOCX, XLSX, CSV and text. "
            "Images need local OCR. Videos are retained for manual review only."
        ),
    )

    docs = []
    for uploaded_file in uploaded or []:
        data = uploaded_file.getvalue()
        pages, metadata, problems = extract_file(uploaded_file.name, data)
        fields, evidence = parse_fields(pages)
        docs.append({
            "name": uploaded_file.name,
            "bytes": data,
            "pages": pages,
            "text": "\n".join(pages),
            "metadata": metadata,
            "problems": problems,
            "fields": fields,
            "evidence": evidence,
            "hash": hashlib.sha256(data).hexdigest(),
        })

    required_docs = RULES.get("categories", {}).get(
        category, {}
    ).get("required_documents", [])

    doc_types = [
        "Application form",
        "Previous EoA / approval",
        "NOC / permission",
        "Accreditation / certification",
        "Faculty records",
        "TER payment proof",
        "Infrastructure photograph",
        "Infrastructure video",
        "Floor plan",
        "Building / safety certificate",
        "Land or building document",
        "Affidavit / undertaking",
        "Other supporting document",
    ]

    if docs:
        st.success(
            f"{len(docs)} file(s) loaded. Assign a type and review extracted "
            "fields before checking."
        )

        for doc in docs:
            with st.expander("📄 " + doc["name"], expanded=True):
                doc["document_type"] = st.selectbox(
                    "Document type",
                    doc_types,
                    key="doctype_" + doc["hash"][:12],
                )

                for field in FIELDS:
                    doc["fields"][field] = st.text_input(
                        field,
                        value=doc["fields"].get(field, ""),
                        key="field_" + doc["hash"][:12] + field,
                    )

                with st.expander(
                    "Show extracted text, file details and processing notes"
                ):
                    st.code(doc["text"][:12000] or "No text extracted.", language=None)
                    st.json({
                        key: value
                        for key, value in doc["metadata"].items()
                        if value
                    })
                    if doc["problems"]:
                        st.warning("\n".join(doc["problems"]))

        if required_docs:
            st.info("Checklist document types: " + ", ".join(required_docs))

        if st.button("Run validation checks", type="primary"):
            findings = run_checks(docs, category, references, role)
            st.session_state["findings"] = findings
            st.session_state["category"] = category
            st.session_state["app_label"] = app_label
            st.session_state["run_id"] = str(uuid.uuid4())
            st.success("Checks complete. Open Dashboard & report for the results.")

    else:
        st.info(
            "Upload documents or evidence to begin. Use synthetic or authorized "
            "anonymized materials for the demo."
        )

    st.caption("Checklist notice: " + RULES.get("notice", "Demo rules only."))

    with st.expander("Infrastructure evidence checklist (prototype aid)"):
        st.write(
            "Attach relevant photos, video, floor plans and certificates where "
            "the chosen application category calls for them. The prototype "
            "records evidence and can extract some text, but it does not measure "
            "buildings, validate engineering compliance, or replace an inspection."
        )
        st.markdown(
            "- Building / floor area evidence\n"
            "- Classrooms, laboratories and equipment\n"
            "- Library and learning resources\n"
            "- Faculty and facilities evidence\n"
            "- Safety, accessibility and other certificates"
        )
        st.caption(
            "This is an illustrative checklist, not an official AICTE checklist. "
            "Confirm applicable requirements from current authorized instructions."
        )

with tab_dashboard:
    findings = st.session_state.get("findings", [])

    if not findings:
        st.info("Upload documents and run checks to populate the dashboard.")
    else:
        passed, flagged, pending = status_counts(findings)
        column1, column2, column3, column4 = st.columns(4)

        column1.metric("Documents checked", len(uploaded or []))
        column2.metric("Checks passed / matched", passed)
        column3.metric("Items to review", flagged)
        column4.metric("Pending / source unavailable", pending)

        st.markdown("#### Verification status")
        st.warning(
            "AICTE live verification: not configured. Synthetic registry matches "
            "are not AICTE confirmations."
        )

        filters = ["All findings"] + sorted({
            finding["status"] for finding in findings
        })
        chosen_filter = st.selectbox("Filter findings", filters)
        shown = (
            findings
            if chosen_filter == "All findings"
            else [
                finding for finding in findings
                if finding["status"] == chosen_filter
            ]
        )

        for finding in shown:
            with st.container(border=True):
                st.markdown(
                    "**" + finding["status"] + "** · " + finding["source"]
                )
                st.write(finding["reason"])
                if finding.get("evidence"):
                    st.caption("Evidence: " + finding["evidence"])

        report = make_report(
            findings,
            st.session_state.get("category", category),
            st.session_state.get("app_label", app_label),
        )
        st.download_button(
            "⬇ Download evidence-linked PDF report",
            report,
            "ai_verify_review_report.pdf",
            "application/pdf",
        )
        st.caption(
            "Extraction accuracy, false alerts, processing time, and source "
            "verification performance need a labeled test set. Do not report "
            "invented scores."
        )

        st.markdown("#### Reviewer decisions")
        for finding in findings:
            with st.expander(finding["status"] + " · " + finding["source"]):
                st.write(finding["reason"])

                with st.form("review_" + finding["id"]):
                    decision = st.selectbox(
                        "Decision",
                        [
                            "Not reviewed",
                            "Resolved",
                            "Needs further verification",
                            "Request additional evidence",
                            "Escalate for inspection",
                        ],
                        key="decision_" + finding["id"],
                    )
                    comment = st.text_input(
                        "Reviewer note",
                        key="comment_" + finding["id"],
                    )

                    if st.form_submit_button("Save review action"):
                        if decision == "Not reviewed":
                            st.warning("Select a decision first.")
                        else:
                            with sqlite3.connect(DB_FILE) as db:
                                db.execute(
                                    "INSERT INTO reviews"
                                    "(finding_id,action,comment,event_time) "
                                    "VALUES(?,?,?,?)",
                                    (
                                        finding["id"],
                                        decision,
                                        comment,
                                        datetime.now().astimezone().isoformat(
                                            timespec="seconds"
                                        ),
                                    ),
                                )

                            log_event(
                                role,
                                "review_action",
                                f"{decision}: {finding['status']}; "
                                f"comment={comment}",
                            )
                            st.success("Review action saved in local audit history.")
                            st.rerun()

with tab_audit:
    st.subheader("Recent local audit events")
    st.caption("This is a local prototype log, not a production secure audit service.")

    events = show_audit()
    if events:
        st.dataframe(
            [
                {
                    "Time": row[0],
                    "Demo role": row[1],
                    "Event": row[2],
                    "Details": row[3],
                }
                for row in events
            ],
            use_container_width=True,
            hide_index=True,
        )
    else:
        st.info("No audit events yet.")

    st.markdown("#### Saved reviewer actions")
    with sqlite3.connect(DB_FILE) as db:
        review_rows = db.execute(
            "SELECT event_time, action, comment "
            "FROM reviews ORDER BY review_id DESC LIMIT 100"
        ).fetchall()

    if review_rows:
        st.dataframe(
            [
                {
                    "Time": row[0],
                    "Action": row[1],
                    "Comment": row[2],
                }
                for row in review_rows
            ],
            use_container_width=True,
            hide_index=True,
        )
    else:
        st.write("No reviewer actions saved yet.")

with st.expander("Issuer / official source verification log"):
    st.caption(
        "Manual prototype log only. No issuing-authority portal, QR service, "
        "signature service or AICTE API is contacted by this form."
    )

    with st.form("source_verification_form"):
        column_a, column_b = st.columns(2)

        source_name = column_a.text_input(
            "Issuing authority / official source"
        )
        source_type = column_b.selectbox(
            "Method",
            [
                "Official portal manual check",
                "QR / reference number manual check",
                "Digital signature manual check",
                "Authorized API (not configured)",
                "Repository lookup manual check",
            ],
        )
        source_reference = st.text_input("Certificate / reference number")
        source_result = st.selectbox(
            "Recorded status",
            [
                "Verification pending",
                "Verified by reviewer",
                "Mismatch detected",
                "Expired",
                "Source unavailable",
                "Manual review required",
            ],
        )
        source_note = st.text_area("Evidence, URL, timestamp or reviewer note")

        if st.form_submit_button("Record source check"):
            if not source_name.strip():
                st.warning("Enter the source name first.")
            else:
                log_event(
                    role,
                    "manual_source_check",
                    json.dumps(
                        {
                            "source": source_name,
                            "method": source_type,
                            "reference": source_reference,
                            "status": source_result,
                            "note": source_note,
                        },
                        ensure_ascii=False,
                    ),
                )
                st.success(
                    "Recorded in the local audit trail. This is a reviewer-entered "
                    "record, not an automated source response."
                )
